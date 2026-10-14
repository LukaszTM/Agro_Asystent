"""Rejestr środków ochrony roślin: XLSX/CSV → SQLite.

Źródło główne: plik „Rejestr zastosowań" ze strony MRiRW (jeden szeroki arkusz,
zwykle świeższy). Uzupełnienie: wydanie kwartalne z dane.gov.pl (licencja CC0)
— stamtąd bierzemy producenta, adres etykiety i słownik upraw.

Wynik: dane/wynik/srodki_ochrony.sqlite oraz raport z liczbami.
Wiersze, których nie da się pewnie odczytać, zapisujemy ze znacznikiem
„do_sprawdzenia" — nigdy z wartością zgadniętą.
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import io
import sqlite3
import sys
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).parent))
from parsery import odczytaj_bbch, odczytaj_dawke, rozbij_uprawy  # noqa: E402

SCHEMAT = """
PRAGMA foreign_keys = ON;

CREATE TABLE zrodlo_danych (
    id             INTEGER PRIMARY KEY AUTOINCREMENT,
    plik           TEXT NOT NULL,
    adres          TEXT,
    licencja       TEXT,
    data_pliku     TEXT,
    data_importu   TEXT NOT NULL,
    liczba_wierszy INTEGER
);

CREATE TABLE srodek (
    id                   INTEGER PRIMARY KEY AUTOINCREMENT,
    nazwa                TEXT NOT NULL,
    nr_zezwolenia        TEXT NOT NULL,
    rodzaj               TEXT,
    substancja_czynna    TEXT,
    producent            TEXT,
    adres_etykiety       TEXT,
    termin_zezwolenia    TEXT,
    termin_do_sprzedazy  TEXT,
    termin_do_stosowania TEXT,
    wycofany             INTEGER NOT NULL DEFAULT 0,
    tylko_profesjonalny  INTEGER,
    mikrobiologiczny     INTEGER,
    agrolotniczy         INTEGER,
    UNIQUE (nazwa, nr_zezwolenia)
);

CREATE TABLE zastosowanie (
    id                INTEGER PRIMARY KEY AUTOINCREMENT,
    srodek_id         INTEGER NOT NULL REFERENCES srodek(id),
    uprawa_surowa     TEXT,
    agrofag           TEXT,
    dawka_tekst       TEXT,
    dawka_wartosc     REAL,
    dawka_jednostka   TEXT,
    dawka_odniesienie TEXT,
    dawka_pewnosc     TEXT NOT NULL,
    termin_tekst      TEXT,
    bbch_od           INTEGER,
    bbch_do           INTEGER,
    grupa             TEXT,
    maloobszarowe     INTEGER,
    metody_stosowania TEXT,
    laczne_stosowanie TEXT,
    -- Wiersz, w którym kolumny rejestru są przesunięte (opis dawki trafił do
    -- kolumny agrofag, a nazwy chorób do kolumny uprawa). Takiego wiersza nie
    -- wolno pokazywać rolnikowi jako zastosowania.
    uklad_podejrzany  INTEGER NOT NULL DEFAULT 0
);

CREATE TABLE zastosowanie_uprawa (
    zastosowanie_id INTEGER NOT NULL REFERENCES zastosowanie(id),
    uprawa          TEXT NOT NULL,
    pewnosc         TEXT NOT NULL CHECK (pewnosc IN ('slownik', 'do_sprawdzenia'))
);

CREATE TABLE uprawa (
    nazwa       TEXT PRIMARY KEY,
    ze_slownika INTEGER NOT NULL
);

CREATE INDEX idx_zastosowanie_srodek ON zastosowanie(srodek_id);
CREATE INDEX idx_zu_uprawa          ON zastosowanie_uprawa(uprawa);
CREATE INDEX idx_zu_zastosowanie    ON zastosowanie_uprawa(zastosowanie_id);
CREATE INDEX idx_srodek_nazwa       ON srodek(nazwa);
"""


#: Opis dawki w kolumnie agrofag zdradza przesunięcie kolumn w wierszu.
WZORZEC_DAWKI_W_AGROFAGU = __import__("re").compile(
    r"maksymalna|zalecana dawka|zastosowania:", __import__("re").IGNORECASE
)


@dataclass
class Raport:
    wierszy_w_pliku: int = 0
    wierszy_z_przesunietymi_kolumnami: int = 0
    srodkow: int = 0
    zastosowan: int = 0
    wycofanych_srodkow: int = 0
    dawka: Counter = field(default_factory=Counter)
    bbch_znane: int = 0
    upraw_rozpoznanych: int = 0
    upraw_do_sprawdzenia: int = 0
    nieznane_uprawy: Counter = field(default_factory=Counter)
    pominiete_wiersze: list[str] = field(default_factory=list)
    uzupelniono_producenta: int = 0
    uzupelniono_etykiete: int = 0


def _tekst(wartosc) -> str | None:
    if wartosc is None:
        return None
    tekst = str(wartosc).strip()
    return tekst or None


def _data(wartosc) -> str | None:
    if isinstance(wartosc, dt.datetime):
        return wartosc.date().isoformat()
    if isinstance(wartosc, dt.date):
        return wartosc.isoformat()
    return _tekst(wartosc)


def _tak_nie(wartosc) -> int | None:
    tekst = _tekst(wartosc)
    if tekst is None:
        return None
    return int(tekst.lower() in {"tak", "t", "1", "true", "x"})


def wczytaj_csv_slownik(sciezka: Path, kolumna: int = 0) -> list[str]:
    """Słowniki z dane.gov.pl bywają w cp1250, nie w UTF-8."""
    surowe = sciezka.read_bytes()
    for kodowanie in ("utf-8-sig", "cp1250", "iso-8859-2"):
        try:
            tekst = surowe.decode(kodowanie)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError(f"Nie rozpoznano kodowania pliku {sciezka}")

    czytnik = csv.reader(io.StringIO(tekst), delimiter=";")
    next(czytnik, None)  # nagłówek
    return [w[kolumna].strip() for w in czytnik if w and w[kolumna].strip()]


def wczytaj_rejestr_podstawowy(sciezka: Path) -> dict[str, dict[str, str | None]]:
    """Producent i adres etykiety w układzie: numer zezwolenia → dane."""
    if not sciezka.exists():
        return {}
    arkusz = openpyxl.load_workbook(sciezka, read_only=True, data_only=True).worksheets[0]
    wiersze = arkusz.iter_rows(values_only=True)
    naglowek = {nazwa: i for i, nazwa in enumerate(next(wiersze)) if nazwa}
    wynik: dict[str, dict[str, str | None]] = {}
    for wiersz in wiersze:
        nr = _tekst(wiersz[naglowek["NrZezw"]]) if "NrZezw" in naglowek else None
        if not nr:
            continue
        wynik[nr] = {
            "producent": _tekst(wiersz[naglowek.get("producent_prosty", -1)]),
            "etykieta": _tekst(wiersz[naglowek.get("etykieta", -1)]),
        }
    return wynik


def importuj(
    plik_zastosowan: Path,
    plik_podstawowy: Path,
    plik_slownika_upraw: Path,
    baza: Path,
    dzis: dt.date,
) -> Raport:
    raport = Raport()

    slownik_upraw = set(wczytaj_csv_slownik(plik_slownika_upraw))
    podstawowy = wczytaj_rejestr_podstawowy(plik_podstawowy)

    arkusz = openpyxl.load_workbook(
        plik_zastosowan, read_only=True, data_only=True
    ).worksheets[0]
    wiersze = arkusz.iter_rows(values_only=True)
    naglowek = {nazwa: i for i, nazwa in enumerate(next(wiersze)) if nazwa}

    wymagane = {"nazwa", "NrZezw", "uprawa", "dawka"}
    brakujace = wymagane - naglowek.keys()
    if brakujace:
        raise ValueError(f"W pliku brakuje kolumn: {sorted(brakujace)}")

    baza.parent.mkdir(parents=True, exist_ok=True)
    if baza.exists():
        baza.unlink()
    polaczenie = sqlite3.connect(baza)
    polaczenie.executescript(SCHEMAT)

    def pole(wiersz, nazwa):
        numer = naglowek.get(nazwa)
        return wiersz[numer] if numer is not None else None

    srodki: dict[tuple[str, str], int] = {}
    uprawy_wszystkie: dict[str, bool] = {}

    for numer_wiersza, wiersz in enumerate(wiersze, start=2):
        if not any(k is not None for k in wiersz):
            continue
        raport.wierszy_w_pliku += 1

        nazwa = _tekst(pole(wiersz, "nazwa"))
        nr_zezw = _tekst(pole(wiersz, "NrZezw"))
        if not nazwa or not nr_zezw:
            raport.pominiete_wiersze.append(
                f"wiersz {numer_wiersza}: brak nazwy środka albo numeru zezwolenia"
            )
            continue

        klucz = (nazwa, nr_zezw)
        if klucz not in srodki:
            termin_do_stosowania = _data(pole(wiersz, "TerminDoStosowania"))
            wycofany = bool(
                termin_do_stosowania and termin_do_stosowania < dzis.isoformat()
            )
            dodatki = podstawowy.get(nr_zezw, {})
            if dodatki.get("producent"):
                raport.uzupelniono_producenta += 1
            if dodatki.get("etykieta"):
                raport.uzupelniono_etykiete += 1
            kursor = polaczenie.execute(
                """INSERT INTO srodek (nazwa, nr_zezwolenia, rodzaj, substancja_czynna,
                       producent, adres_etykiety, termin_zezwolenia, termin_do_sprzedazy,
                       termin_do_stosowania, wycofany, tylko_profesjonalny,
                       mikrobiologiczny, agrolotniczy)
                   VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (
                    nazwa,
                    nr_zezw,
                    _tekst(pole(wiersz, "Rodzaj")),
                    _tekst(pole(wiersz, "Substancja_czynna")),
                    dodatki.get("producent"),
                    dodatki.get("etykieta"),
                    _data(pole(wiersz, "TerminZezw")),
                    _data(pole(wiersz, "TerminDoSprzedazy")),
                    termin_do_stosowania,
                    int(wycofany),
                    int((_tekst(pole(wiersz, "zastosowanie/uzytkownik")) or "").lower()
                        == "profesjonalne"),
                    _tak_nie(pole(wiersz, "srodek_mikrobiologiczny")),
                    _tak_nie(pole(wiersz, "stosowanie_agrolotnicze")),
                ),
            )
            srodki[klucz] = int(kursor.lastrowid)
            raport.srodkow += 1
            raport.wycofanych_srodkow += int(wycofany)

        agrofag = _tekst(pole(wiersz, "agrofag"))
        uklad_podejrzany = bool(agrofag and WZORZEC_DAWKI_W_AGROFAGU.search(agrofag))
        if uklad_podejrzany:
            raport.wierszy_z_przesunietymi_kolumnami += 1

        dawka = odczytaj_dawke(_tekst(pole(wiersz, "dawka")))
        bbch_od, bbch_do = odczytaj_bbch(_tekst(pole(wiersz, "termin")))
        raport.dawka[dawka.pewnosc] += 1
        raport.bbch_znane += int(bbch_od is not None)

        kursor = polaczenie.execute(
            """INSERT INTO zastosowanie (srodek_id, uprawa_surowa, agrofag, dawka_tekst,
                   dawka_wartosc, dawka_jednostka, dawka_odniesienie, dawka_pewnosc,
                   termin_tekst, bbch_od, bbch_do, grupa, maloobszarowe,
                   metody_stosowania, laczne_stosowanie, uklad_podejrzany)
               VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (
                srodki[klucz],
                _tekst(pole(wiersz, "uprawa")),
                agrofag,
                dawka.tekst or None,
                dawka.wartosc,
                dawka.jednostka,
                dawka.odniesienie,
                dawka.pewnosc,
                _tekst(pole(wiersz, "termin")),
                bbch_od,
                bbch_do,
                _tekst(pole(wiersz, "nazwa_grupy")) or _tekst(pole(wiersz, "grupa")),
                _tak_nie(pole(wiersz, "maloobszarowe")),
                _tekst(pole(wiersz, "metody_stosowania")),
                _tekst(pole(wiersz, "laczne_stosowanie")),
                int(uklad_podejrzany),
            ),
        )
        id_zastosowania = int(kursor.lastrowid)
        raport.zastosowan += 1

        if uklad_podejrzany:
            # W takim wierszu kolumna „uprawa" zawiera nazwy chorób — nie
            # rozbijamy jej na uprawy, bo powstałyby fałszywe pozycje.
            continue

        rozbite = rozbij_uprawy(_tekst(pole(wiersz, "uprawa")), slownik_upraw)
        for nazwa_uprawy in rozbite.rozpoznane:
            polaczenie.execute(
                "INSERT INTO zastosowanie_uprawa VALUES (?,?,?)",
                (id_zastosowania, nazwa_uprawy, "slownik"),
            )
            uprawy_wszystkie[nazwa_uprawy] = True
            raport.upraw_rozpoznanych += 1
        for nazwa_uprawy in rozbite.do_sprawdzenia:
            polaczenie.execute(
                "INSERT INTO zastosowanie_uprawa VALUES (?,?,?)",
                (id_zastosowania, nazwa_uprawy, "do_sprawdzenia"),
            )
            uprawy_wszystkie.setdefault(nazwa_uprawy, False)
            raport.upraw_do_sprawdzenia += 1
            raport.nieznane_uprawy[nazwa_uprawy] += 1

    polaczenie.executemany(
        "INSERT INTO uprawa VALUES (?,?)",
        [(n, int(z)) for n, z in sorted(uprawy_wszystkie.items())],
    )
    polaczenie.execute(
        """INSERT INTO zrodlo_danych (plik, adres, licencja, data_pliku,
               data_importu, liczba_wierszy) VALUES (?,?,?,?,?,?)""",
        (
            plik_zastosowan.name,
            "https://www.gov.pl/web/rolnictwo/wyszukiwarka-srodkow-ochrony-roslin",
            "materiał urzędowy MRiRW",
            dt.datetime.fromtimestamp(plik_zastosowan.stat().st_mtime).date().isoformat(),
            dzis.isoformat(),
            raport.wierszy_w_pliku,
        ),
    )
    if plik_slownika_upraw.exists():
        polaczenie.execute(
            """INSERT INTO zrodlo_danych (plik, adres, licencja, data_pliku,
                   data_importu, liczba_wierszy) VALUES (?,?,?,?,?,?)""",
            (
                plik_slownika_upraw.name,
                "https://dane.gov.pl — Wyszukiwarka środków ochrony roślin",
                "CC0 1.0",
                dt.datetime.fromtimestamp(plik_slownika_upraw.stat().st_mtime)
                .date()
                .isoformat(),
                dzis.isoformat(),
                len(slownik_upraw),
            ),
        )
    polaczenie.commit()
    polaczenie.close()
    return raport


def wypisz_raport(raport: Raport, baza: Path) -> None:
    razem = max(raport.zastosowan, 1)
    upraw_razem = max(raport.upraw_rozpoznanych + raport.upraw_do_sprawdzenia, 1)
    print(f"\nBaza: {baza}  ({baza.stat().st_size / 1_048_576:.1f} MB)")
    print(f"  wierszy w pliku:        {raport.wierszy_w_pliku}")
    print(f"  środków:                {raport.srodkow} "
          f"(po terminie stosowania: {raport.wycofanych_srodkow})")
    print(f"  zastosowań:             {raport.zastosowan}")
    print(f"  wierszy z przesuniętymi kolumnami: "
          f"{raport.wierszy_z_przesunietymi_kolumnami} (oznaczone, bez rozbicia upraw)")
    print(f"  uzupełniono producenta: {raport.uzupelniono_producenta}")
    print(f"  uzupełniono etykietę:   {raport.uzupelniono_etykiete}")
    print("\n  dawka:")
    for pewnosc in ("odczytana", "zakres", "do_sprawdzenia"):
        ile = raport.dawka.get(pewnosc, 0)
        print(f"    {pewnosc:16} {ile:6}  ({100 * ile / razem:5.1f}%)")
    print(f"\n  faza BBCH odczytana:    {raport.bbch_znane} "
          f"({100 * raport.bbch_znane / razem:.1f}%)")
    print(f"  uprawy ze słownika:     {raport.upraw_rozpoznanych} "
          f"({100 * raport.upraw_rozpoznanych / upraw_razem:.1f}%)")
    print(f"  uprawy do sprawdzenia:  {raport.upraw_do_sprawdzenia} "
          f"({len(raport.nieznane_uprawy)} różnych nazw)")
    if raport.nieznane_uprawy:
        print("\n  najczęstsze nazwy upraw spoza słownika:")
        for nazwa, ile in raport.nieznane_uprawy.most_common(15):
            print(f"    {nazwa[:52]:52} {ile}")
    if raport.pominiete_wiersze:
        print(f"\n  pominięte wiersze: {len(raport.pominiete_wiersze)}")
        for opis in raport.pominiete_wiersze[:10]:
            print(f"    {opis}")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Importuje rejestr środków ochrony roślin do bazy SQLite."
    )
    parser.add_argument("--katalog", type=Path, default=Path("dane/pobrane/p1"),
                        help="katalog z pobranymi plikami")
    parser.add_argument("--zastosowania", type=Path, default=None,
                        help="plik z zastosowaniami (domyślnie MRiRW z --katalog)")
    parser.add_argument("--baza", type=Path,
                        default=Path("dane/wynik/srodki_ochrony.sqlite"),
                        help="plik bazy do zapisania")
    args = parser.parse_args(argv)

    zastosowania = args.zastosowania or args.katalog / "MRiRW_Rejestr_zastosowanie.xlsx"
    if not zastosowania.exists():
        print(f"Brak pliku {zastosowania}. Uruchom najpierw pobierz.py", file=sys.stderr)
        return 1

    raport = importuj(
        plik_zastosowan=zastosowania,
        plik_podstawowy=args.katalog / "Rejestr_podstawowe.xlsx",
        plik_slownika_upraw=args.katalog / "Slownik_Uprawy.csv",
        baza=args.baza,
        dzis=dt.date.today(),
    )
    wypisz_raport(raport, args.baza)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
